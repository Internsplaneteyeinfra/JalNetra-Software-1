"""Per-scene flood/water time series with rolling pre/post scene pairs."""
from __future__ import annotations

import io
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import ee
import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import plotly.graph_objects as go
from openpyxl import Workbook

from flood_service import (
    _SCALE,
    _build_classification,
    _compute_flood_water_areas_ha,
    _filtered_collection,
    _next_day,
    _s1_collection,
    _tile_url,
)
from speckle_filters import refined_lee, to_db, to_natural

_ts_executor = ThreadPoolExecutor(max_workers=8, thread_name_prefix="flood_ts_ee_")
_EXCEL_SCALE = _SCALE  # same as analysis (10 m) so each water/flood pixel is exported
_EXCEL_CHUNK = 4500  # stay under EE's 5000-element getInfo limit
_EXCEL_MAX_POINTS = 500000  # per-scene safety cap for very large inundation areas
# Monsoon months only for Excel export (June–September)
_EXCEL_MONTHS = frozenset({6, 7, 8, 9})
_DEFAULT_YEAR_RANGE = (2015, 2026)


def _geometry_centroid_latlon(geometry: ee.Geometry) -> Tuple[float, float]:
    coords = geometry.centroid(maxError=1).coordinates().getInfo()
    lon = float(coords[0]) if coords else 0.0
    lat = float(coords[1]) if coords else 0.0
    return round(lat, 6), round(lon, 6)


def _scene_records(
    collection: ee.ImageCollection, geometry: ee.Geometry, start_date: str, end_date: str
) -> List[Dict[str, Any]]:
    ee_end = _next_day(end_date)
    filtered = (
        _filtered_collection(collection, start_date, ee_end, geometry)
        .sort("system:time_start")
    )
    count = int(filtered.size().getInfo() or 0)
    if count == 0:
        return []

    def _feature(img: ee.Image) -> ee.Feature:
        millis = img.get("system:time_start")
        return ee.Feature(
            None,
            {
                "system:time_start": millis,
                "date": ee.Date(millis).format("YYYY-MM-dd"),
            },
        )

    features = filtered.map(_feature).getInfo().get("features", [])
    records: List[Dict[str, Any]] = []
    for feat in features:
        props = feat.get("properties") or {}
        millis = props.get("system:time_start")
        date_str = props.get("date")
        if millis is None or not date_str:
            continue
        records.append(
            {
                "date": date_str,
                "millis": int(millis),
                "month": date_str[:7],
            }
        )
    records.sort(key=lambda r: r["millis"])
    return records


def _image_for_millis(
    collection: ee.ImageCollection, geometry: ee.Geometry, millis: int
) -> ee.Image:
    img = (
        collection.filter(ee.Filter.eq("system:time_start", millis))
        .first()
        .select("VH")
        .clip(geometry)
    )
    return img


def _last_scene_before_month(
    collection: ee.ImageCollection, geometry: ee.Geometry, month_key: str
) -> Optional[Tuple[ee.Image, str]]:
    """Last Sentinel-1 scene strictly before the first day of month_key (YYYY-MM)."""
    month_start = f"{month_key}-01"
    filtered = (
        collection.filter(ee.Filter.date("1970-01-01", month_start))
        .filterBounds(geometry)
        .sort("system:time_start", False)
    )
    if int(filtered.size().getInfo() or 0) == 0:
        return None
    img = ee.Image(filtered.first()).select("VH").clip(geometry)
    pre_date = (
        ee.Date(filtered.first().get("system:time_start"))
        .format("YYYY-MM-dd")
        .getInfo()
    )
    return img, pre_date


def _extract_classified_points(
    classification: ee.Image,
    geometry: ee.Geometry,
    post_date: str,
    pre_date: str,
    scale: int = _EXCEL_SCALE,
) -> List[Dict[str, Any]]:
    """Export lat/lon for every water/flood pixel (paginated under EE 5000 limit)."""
    sampled_img = (
        classification.rename("class")
        .addBands(ee.Image.pixelLonLat())
        .updateMask(classification.gt(0))
    )
    fc = sampled_img.sample(
        region=geometry,
        scale=scale,
        geometries=False,
        tileScale=16,
        dropNulls=True,
    )

    total = int(fc.size().getInfo() or 0)
    if total == 0:
        return []

    limit = min(total, _EXCEL_MAX_POINTS)
    rows: List[Dict[str, Any]] = []
    for offset in range(0, limit, _EXCEL_CHUNK):
        chunk_n = min(_EXCEL_CHUNK, limit - offset)
        features = (
            ee.FeatureCollection(fc.toList(chunk_n, offset))
            .getInfo()
            .get("features", [])
        )
        for feat in features:
            props = feat.get("properties") or {}
            cls = int(float(props.get("class") or 0))
            if cls not in (1, 2):
                continue
            lat = props.get("latitude")
            lon = props.get("longitude")
            if lat is None or lon is None:
                continue
            rows.append(
                {
                    "post_date": str(post_date),
                    "pre_date": str(pre_date),
                    "latitude": round(float(lat), 6),
                    "longitude": round(float(lon), 6),
                    "class": "water" if cls == 1 else "flood",
                }
            )
    return rows


def _month_number(date_str: str) -> int:
    """Extract calendar month (1–12) from YYYY-MM-DD."""
    return int(date_str[5:7])


def _year_number(date_str: str) -> int:
    """Extract calendar year from YYYY-MM-DD."""
    return int(date_str[:4])


def build_timeseries_excel(
    comparisons: List[Dict[str, Any]],
    scene_points: Dict[str, List[Dict[str, Any]]],
) -> bytes:
    """
    Exactly two sheets (all months/years stacked in date order):
      1) areas  — pre_date, post_date, water_area_ha, flood_area_ha
      2) lat_lon — pre_date, post_date, latitude, longitude, class
    """
    wb = Workbook(write_only=True)
    area_headers = ["pre_date", "post_date", "water_area_ha", "flood_area_ha"]
    point_headers = ["pre_date", "post_date", "latitude", "longitude", "class"]

    areas = wb.create_sheet(title="areas")
    areas.append(area_headers)
    for comp in sorted(comparisons, key=lambda c: c.get("post_date") or ""):
        areas.append(
            [
                str(comp.get("pre_date") or ""),
                str(comp.get("post_date") or ""),
                float(comp.get("water_area_ha") or 0),
                float(comp.get("flood_area_ha") or 0),
            ]
        )

    lat_lon = wb.create_sheet(title="lat_lon")
    lat_lon.append(point_headers)
    for post_date in sorted(scene_points.keys()):
        for row in scene_points.get(post_date) or []:
            lat_lon.append(
                [
                    str(row.get("pre_date") or ""),
                    str(row.get("post_date") or post_date),
                    float(row.get("latitude") or 0),
                    float(row.get("longitude") or 0),
                    str(row.get("class") or ""),
                ]
            )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _analyze_scene_pair(
    pre_image: ee.Image,
    post_image: ee.Image,
    geometry: ee.Geometry,
    pre_date: str,
    post_date: str,
    post_millis: int,
    month: str,
    *,
    include_scene_tiles: bool,
    include_excel_points: bool,
) -> Dict[str, Any]:
    before_filtered = ee.Image(to_db(refined_lee(to_natural(pre_image))))
    after_filtered = ee.Image(to_db(refined_lee(to_natural(post_image))))
    classification, water_mask, flood_mask = _build_classification(
        before_filtered, after_filtered, geometry
    )

    water_ha, flood_ha = _compute_flood_water_areas_ha(water_mask, flood_mask, geometry)
    result: Dict[str, Any] = {
        "month": month,
        "post_date": post_date,
        "pre_date": pre_date,
        "post_millis": post_millis,
        "flood_area_ha": flood_ha,
        "water_area_ha": water_ha,
        "total_water_available_ha": water_ha,
    }
    if include_scene_tiles:
        result["tile_url"] = _tile_url(classification, geometry)
    if include_excel_points:
        try:
            result["point_rows"] = _extract_classified_points(
                classification,
                geometry,
                post_date,
                pre_date,
            )
        except Exception as exc:
            result["point_rows"] = []
            result["excel_points_error"] = str(exc)
    return result


def _build_comparison_pairs(
    collection: ee.ImageCollection,
    geometry: ee.Geometry,
    scenes: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    by_month: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for scene in scenes:
        by_month[scene["month"]].append(scene)

    pairs: List[Dict[str, Any]] = []
    for month in sorted(by_month.keys()):
        month_scenes = by_month[month]
        for idx, post_scene in enumerate(month_scenes):
            if idx == 0:
                pre_lookup = _last_scene_before_month(collection, geometry, month)
                if pre_lookup is None:
                    continue
                pre_image, pre_date = pre_lookup
            else:
                pre_scene = month_scenes[idx - 1]
                pre_image = _image_for_millis(collection, geometry, pre_scene["millis"])
                pre_date = pre_scene["date"]

            post_image = _image_for_millis(collection, geometry, post_scene["millis"])
            pairs.append(
                {
                    "pre_image": pre_image,
                    "post_image": post_image,
                    "pre_date": pre_date,
                    "post_date": post_scene["date"],
                    "post_millis": post_scene["millis"],
                    "month": month,
                }
            )
    return pairs


def _build_timeseries_figure(comparisons: List[Dict[str, Any]]) -> go.Figure:
    if not comparisons:
        raise ValueError("No comparisons available to build graph.")

    dates = [c["post_date"] for c in comparisons]
    flood_vals = [c["flood_area_ha"] for c in comparisons]
    water_vals = [c["water_area_ha"] for c in comparisons]
    pre_dates = [c["pre_date"] for c in comparisons]

    fig = go.Figure()
    fig.add_trace(
        go.Scatter(
            x=dates,
            y=flood_vals,
            mode="lines+markers",
            name="Flood",
            marker=dict(size=13, color="#ff0000", symbol="circle", line=dict(width=2, color="#ffffff")),
            line=dict(color="#ff0000", width=2, dash="solid"),
            customdata=pre_dates,
            hovertemplate=(
                "<b>Flood</b><br>"
                "Date: %{x}<br>"
                "Pre date: %{customdata}<br>"
                "Area: %{y:.2f} ha"
                "<extra></extra>"
            ),
        )
    )
    fig.add_trace(
        go.Scatter(
            x=dates,
            y=water_vals,
            mode="lines+markers",
            name="Water",
            marker=dict(size=13, color="#0000ff", symbol="circle", line=dict(width=2, color="#ffffff")),
            line=dict(color="#0000ff", width=2, dash="solid"),
            customdata=pre_dates,
            hovertemplate=(
                "<b>Water</b><br>"
                "Date: %{x}<br>"
                "Pre date: %{customdata}<br>"
                "Area: %{y:.2f} ha"
                "<extra></extra>"
            ),
        )
    )
    fig.update_layout(
        title="June–September Flood & Water Time Series (2015–2026)",
        xaxis_title="Date",
        yaxis_title="water and flood Area (hectares)",
        template="plotly_white",
        hovermode="closest",
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
        height=520,
        width=max(900, len(dates) * 72),
        margin=dict(l=80, r=30, t=70, b=80),
    )
    fig.update_xaxes(showgrid=True, gridcolor="#e5e7eb", tickangle=-45)
    fig.update_yaxes(showgrid=True, gridcolor="#e5e7eb")
    return fig


def _render_timeseries_html(comparisons: List[Dict[str, Any]]) -> bytes:
    fig = _build_timeseries_figure(comparisons)
    html = fig.to_html(
        include_plotlyjs="cdn",
        full_html=True,
        config={"displayModeBar": True, "responsive": True},
        div_id="flood-water-chart",
    )

    click_panel = """
<div style="margin-top:14px;display:flex;gap:10px;flex-wrap:wrap;align-items:center;">
  <button id="download-png" type="button" style="padding:8px 14px;border:none;border-radius:6px;background:#374151;color:#fff;font-size:14px;cursor:pointer;">
    Download PNG
  </button>
  <div id="point-info" style="flex:1;min-width:240px;padding:12px 14px;border:1px solid #d1d5db;border-radius:8px;background:#f9fafb;font-family:system-ui,sans-serif;font-size:14px;">
    Hover or click a dot to see type, date, and area (ha).
  </div>
</div>
<script>
  const chartEl = document.getElementById('flood-water-chart');
  function showPoint(pt) {
    const type = pt.data.name || 'Point';
    const date = pt.x;
    const area = Number(pt.y).toFixed(2);
    const pre = pt.customdata || '-';
    document.getElementById('point-info').innerHTML =
      '<strong>' + type + '</strong><br>' +
      'Date: ' + date + '<br>' +
      'Pre date: ' + pre + '<br>' +
      'Area: ' + area + ' ha';
  }
  chartEl.on('plotly_hover', function(e) { if (e.points && e.points[0]) showPoint(e.points[0]); });
  chartEl.on('plotly_click', function(e) { if (e.points && e.points[0]) showPoint(e.points[0]); });
  document.getElementById('download-png').addEventListener('click', function() {
    const btn = this;
    btn.disabled = true;
    btn.textContent = 'Preparing PNG...';
    Plotly.downloadImage(chartEl, {
      format: 'png',
      width: chartEl.layout.width || 1100,
      height: chartEl.layout.height || 520,
      filename: 'flood_water_timeseries'
    }).then(function() {
      btn.disabled = false;
      btn.textContent = 'Download PNG';
    }).catch(function() {
      btn.disabled = false;
      btn.textContent = 'Download PNG';
      alert('PNG download failed in browser. Use the Download PNG link on the graph viewer page.');
    });
  });
</script>
"""
    if "</body>" in html:
        html = html.replace("</body>", click_panel + "</body>")
    else:
        html += click_panel

    return html.encode("utf-8")


def render_timeseries_png(comparisons: List[Dict[str, Any]]) -> bytes:
    """Export PNG via matplotlib (kaleido hangs on many Windows setups)."""
    dates = [datetime.strptime(c["post_date"], "%Y-%m-%d") for c in comparisons]
    flood_vals = [c["flood_area_ha"] for c in comparisons]
    water_vals = [c["water_area_ha"] for c in comparisons]

    fig_width = max(11.0, len(dates) * 1.1)
    fig, ax = plt.subplots(figsize=(fig_width, 5.2))
    ax.plot(
        dates,
        flood_vals,
        "o-",
        color="#ff0000",
        linewidth=2,
        markersize=9,
        markerfacecolor="#ff0000",
        markeredgecolor="#ffffff",
        markeredgewidth=2,
        label="Flood",
    )
    ax.plot(
        dates,
        water_vals,
        "o-",
        color="#0000ff",
        linewidth=2,
        markersize=9,
        markerfacecolor="#0000ff",
        markeredgecolor="#ffffff",
        markeredgewidth=2,
        label="Water",
    )
    ax.set_title("Flood & Water by Sentinel-1 Scene (connected dots)")
    ax.set_xlabel("Date")
    ax.set_ylabel("Area (hectares)")
    ax.grid(True, color="#e5e7eb", linestyle="-", linewidth=0.8)
    ax.legend(loc="upper left")
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y-%m-%d"))
    fig.autofmt_xdate(rotation=45, ha="right")
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


def _run_scene_comparisons(
    geometry: ee.Geometry,
    start_date: str,
    end_date: str,
    *,
    include_scene_tiles: bool = False,
    include_excel_points: bool = False,
    months_filter: Optional[frozenset] = None,
    year_range: Optional[Tuple[int, int]] = None,
) -> Dict[str, Any]:
    """Run rolling pre/post scene comparisons for [start_date, end_date]."""
    collection = _s1_collection()
    scenes = _scene_records(collection, geometry, start_date, end_date)
    if months_filter:
        scenes = [
            s for s in scenes if _month_number(s["date"]) in months_filter
        ]
    if year_range:
        y0, y1 = year_range
        scenes = [s for s in scenes if y0 <= _year_number(s["date"]) <= y1]
    if not scenes:
        return {
            "scenes": [],
            "comparisons": [],
            "message": "No Sentinel-1 scenes found in the selected date ranges.",
        }

    pair_specs = _build_comparison_pairs(collection, geometry, scenes)
    if not pair_specs:
        return {
            "scenes": scenes,
            "comparisons": [],
            "message": "No pre/post pairs could be built (missing prior-month scene).",
        }

    futures = [
        _ts_executor.submit(
            _analyze_scene_pair,
            spec["pre_image"],
            spec["post_image"],
            geometry,
            spec["pre_date"],
            spec["post_date"],
            spec["post_millis"],
            spec["month"],
            include_scene_tiles=include_scene_tiles,
            include_excel_points=include_excel_points,
        )
        for spec in pair_specs
    ]
    comparisons = [f.result() for f in futures]
    comparisons.sort(key=lambda c: c["post_millis"])
    return {"scenes": scenes, "comparisons": comparisons, "message": None}


def build_scene_timeseries(
    geometry: ee.Geometry,
    pre_date: str,
    post_date: str,
    *,
    include_scene_tiles: bool = False,
    include_excel: bool = True,
) -> Dict[str, Any]:
    """
    Graph and Excel use pre_date → post_date, filtered to June–September (2015–2026).

    Excel (when enabled) writes exactly two sheets with rows stacked
    chronologically by month/year:
      - areas
      - lat_lon
    """
    run = _run_scene_comparisons(
        geometry,
        pre_date,
        post_date,
        include_scene_tiles=include_scene_tiles,
        include_excel_points=include_excel,
        months_filter=_EXCEL_MONTHS,
        year_range=_DEFAULT_YEAR_RANGE,
    )
    comparisons = run["comparisons"]

    excel_bytes = None
    excel_point_errors: List[str] = []
    excel_range = {"start": pre_date, "end": post_date} if include_excel else None
    scene_points: Dict[str, List[Dict[str, Any]]] = {}
    if include_excel:
        for comp in comparisons:
            err = comp.pop("excel_points_error", None)
            if err:
                excel_point_errors.append(f"{comp.get('post_date')}: {err}")
            rows = comp.pop("point_rows", []) or []
            if rows:
                scene_points[str(comp["post_date"])] = rows
        excel_bytes = build_timeseries_excel(comparisons, scene_points)
    else:
        for comp in comparisons:
            comp.pop("point_rows", None)
            comp.pop("excel_points_error", None)

    graph_html = (
        _render_timeseries_html(comparisons) if comparisons else None
    )

    result_payload: Dict[str, Any] = {
        "scene_count": len(run["scenes"]),
        "comparison_count": len(comparisons),
        "comparisons": comparisons,
        "graph_html": graph_html,
        "excel_bytes": excel_bytes,
        "excel_date_range": excel_range,
        "message": run.get("message"),
        "logic": (
            "First scene each month uses the previous month's last image as pre; "
            "each following scene uses the prior scene in the same month as pre. "
            "Excel has two sheets (areas + lat_lon) with all months/years "
            "stacked in date order."
        ),
    }
    if excel_point_errors:
        result_payload["excel_points_warning"] = (
            "Some scene lat/lon extractions failed: "
            + "; ".join(excel_point_errors[:5])
        )
    return result_payload